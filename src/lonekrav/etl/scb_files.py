"""SCB SSYK 2012 workbooks: structure, English titles, title index, ISCO-08 key."""

from __future__ import annotations

import io
from collections import defaultdict
from pathlib import Path

import openpyxl

from lonekrav.etl.http import cached_get

BASE = "https://www.scb.se/contentassets/0c0089cc085a45d49c1dc83923ad933a/"
FILES = {
    "structure": "ssyk-2012-koder.xlsx",
    "english": "structur_english_ssyk2012_1_4.xlsx",
    "title_index": "ssyk_index_webb_2026.xlsx",
    "isco_key": "webb_nyckel_ssyk2012_isco-08_20160905.xlsx",
}
SHEETS = {"structure": "4-siffer", "english": "4-digit", "title_index": "Index",
          "isco_key": "Nyckel"}


def validator(key: str):
    """Reject a download that is not a workbook with the sheet its parser reads."""
    def check(data: bytes) -> None:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        if SHEETS[key] not in wb.sheetnames:
            raise ValueError(f"no sheet '{SHEETS[key]}' (found {wb.sheetnames})")
    return check


def fetch(raw: Path, force: bool = False) -> dict[str, Path]:
    return {k: cached_get("scb_files", BASE + f, raw / "scb" / f, force=force,
                          validate=validator(k))
            for k, f in FILES.items()}


def _code(v) -> str | None:
    """Four-digit SSYK code, padding codes stored as numbers ('110' -> '0110')."""
    s = str(v).strip() if v is not None else ""
    return s.zfill(4) if s.isdigit() and len(s) in (3, 4) else None


def _rows(path: Path, sheet: str):
    wb = openpyxl.load_workbook(path, read_only=True)
    yield from wb[sheet].iter_rows(values_only=True)


def parse_structure(path: Path) -> dict[str, str]:
    return {c: str(r[1]).strip() for r in _rows(path, "4-siffer")
            if (c := _code(r[0])) and r[1] and len(str(r[0]).strip()) == 4}


def parse_english(path: Path) -> dict[str, str]:
    return {c: str(r[1]).strip() for r in _rows(path, "4-digit")
            if (c := _code(r[0])) and r[1] and len(str(r[0]).strip()) == 4}


def parse_major_groups(structure: Path, english: Path) -> dict[str, dict[str, str]]:
    """SSYK's ten major groups (the first digit), in Swedish and English, for the occupation map."""
    def one(path: Path, sheet: str) -> dict[str, str]:
        return {str(r[0]).strip(): str(r[1]).strip() for r in _rows(path, sheet)
                if r[0] is not None and r[1] and str(r[0]).strip().isdigit() and len(str(r[0]).strip()) == 1}
    sv, en = one(structure, "1-siffer"), one(english, "1-digit")
    return {d: {"sv": sv[d], "en": en[d]} for d in sorted(sv) if d in en}  # all ten in SCB's files


def parse_isco_key(path: Path) -> dict[str, list[str]]:
    key = {}
    for r in _rows(path, "Nyckel"):
        if (c := _code(r[0])) and r[2] is not None:
            key[c] = [i.strip().zfill(4) for i in str(r[2]).split(",") if i.strip()]
    return key


def parse_title_index(path: Path) -> dict[str, list[str]]:
    idx: dict[str, list[str]] = defaultdict(list)
    for r in _rows(path, "Index"):
        if (c := _code(r[0])) and r[1]:
            idx[c].append(str(r[1]).strip())
    return dict(idx)
