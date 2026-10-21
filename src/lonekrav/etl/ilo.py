"""ILO ISCO-08 English structure and definitions (unit groups only)."""

from __future__ import annotations

import io
import re
from pathlib import Path

import openpyxl
from pydantic import BaseModel

from lonekrav.etl.http import cached_get

URL = ("https://www.ilo.org/ilostat-files/ISCO/newdocs-08-2021/ISCO-08/"
       "ISCO-08%20EN%20Structure%20and%20definitions.xlsx")


class UnitGroup(BaseModel):
    code: str
    title: str
    definition: str
    tasks: str
    included: list[str]


def validate(data: bytes) -> None:
    header = next(openpyxl.load_workbook(io.BytesIO(data), read_only=True).worksheets[0]
                  .iter_rows(max_row=1, values_only=True))
    if "ISCO 08 Code" not in header:
        raise ValueError(f"unexpected header {header}")


def fetch(raw: Path, force: bool = False) -> Path:
    return cached_get("ilo", URL, raw / "ilo" / "isco08_en.xlsx", force=force, validate=validate)


def _text(v) -> str:
    return str(v).strip() if v is not None else ""


def _bullets(v) -> list[str]:
    """'Examples of the occupations classified here:\\n- A\\n- B' -> ['A', 'B']."""
    return [m.strip() for m in re.findall(r"(?m)^\s*-\s*(.+?)\s*$", _text(v))]


def parse(path: Path) -> dict[str, UnitGroup]:
    groups = {}
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if _text(r[0]) != "4":
            continue
        code = _text(r[1]).zfill(4)
        tasks = re.sub(r"^Tasks include\s*-?\s*", "", _text(r[4]))
        groups[code] = UnitGroup(code=code, title=_text(r[2]), definition=_text(r[3]),
                                 tasks=tasks, included=_bullets(r[5]))
    return groups
