"""Download each Lönekrav source once and keep small, real slices as test fixtures.

Standalone on purpose (plain httpx + openpyxl + pypdf), so parser tests have real inputs before
any parser exists. Run from the repo root:

    uv run python scripts/cut_fixtures.py

Writes to tests/fixtures/ and records each URL with its fetch date in tests/fixtures/README.md.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import re
from pathlib import Path

import httpx
import openpyxl
from pypdf import PdfReader

OUT = Path(__file__).resolve().parents[1] / "tests" / "fixtures"
CODES = ["0000", "0110", "1211", "1212", "2421", "2512", "3511", "4430", "5322", "5343", "9210"]
SCB_FILES = "https://www.scb.se/contentassets/0c0089cc085a45d49c1dc83923ad933a/"
WORKBOOKS = {
    "ssyk-2012-koder.xlsx": SCB_FILES + "ssyk-2012-koder.xlsx",
    "structur_english_ssyk2012_1_4.xlsx": SCB_FILES + "structur_english_ssyk2012_1_4.xlsx",
    "ssyk_index_webb_2026.xlsx": SCB_FILES + "ssyk_index_webb_2026.xlsx",
    "webb_nyckel_ssyk2012_isco-08_20160905.xlsx": SCB_FILES
    + "webb_nyckel_ssyk2012_isco-08_20160905.xlsx",
}
PXWEB = "https://statistikdatabasen.scb.se/api/v2/tables/{table}/data"
PXWEB_TABLES = {  # table -> extra selections
    "TAB5932": {},
    "TAB4265": {},
    "TAB5844": {"Alder": "*"},
    "TAB4264": {"Alder": "*"},
}
JOBTECH = "https://taxonomy.api.jobtechdev.se/v1/taxonomy/graphql"
JOBTECH_QUERY = """{ concepts(type:"ssyk-level-4") { id ssyk_code_2012 preferred_label definition
  narrower(type:"occupation-name") { id preferred_label alternative_labels
    exact_match(type:"esco-occupation") { esco_uri preferred_label }
    close_match(type:"esco-occupation") { esco_uri preferred_label }
    broad_match(type:"esco-occupation") { esco_uri preferred_label }
    narrow_match(type:"esco-occupation") { esco_uri preferred_label } } } }"""
ILO_URL = (
    "https://www.ilo.org/ilostat-files/ISCO/newdocs-08-2021/ISCO-08/"
    "ISCO-08%20EN%20Structure%20and%20definitions.xlsx"
)
ESCO_URL = "https://ec.europa.eu/esco/api/resource/occupation"
MIS_URL = (
    "https://www.scb.se/contentassets/c9d055b6f2114b62bd23c33602b56da5/"
    "ov9999_2012a01_br_x70br1201.pdf"
)

client = httpx.Client(timeout=120, follow_redirects=True)
sources: list[tuple[str, str]] = []


def get(url: str, **kw) -> httpx.Response:
    r = client.get(url, **kw)
    r.raise_for_status()
    sources.append((url, dt.date.today().isoformat()))
    return r


def keep_rows(src: bytes, dest: Path, keep) -> None:
    """Re-save a workbook with only the header rows and the rows `keep(row)` accepts."""
    wb = openpyxl.load_workbook(io.BytesIO(src), read_only=True)
    out = openpyxl.Workbook()
    out.remove(out.active)
    for ws in wb.worksheets:
        o = out.create_sheet(ws.title)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 5 or keep(row):
                o.append(list(row))
    out.save(dest)


def code_of(row) -> str:
    v = row[0] if row else None
    return str(v).strip().zfill(4) if v is not None and str(v).strip().isdigit() else ""


def workbooks() -> set[str]:
    isco: set[str] = set()
    for name, url in WORKBOOKS.items():
        src = get(url).content
        keep_rows(src, OUT / "scb" / name, lambda r: code_of(r) in CODES)
        if name.startswith("webb_nyckel"):
            wb = openpyxl.load_workbook(io.BytesIO(src), read_only=True)
            for r in wb["Nyckel"].iter_rows(values_only=True):
                if code_of(r) in CODES and r[2] is not None:
                    isco |= {c.strip().zfill(4) for c in str(r[2]).split(",")}
    return isco


def pxweb() -> None:
    for table, extra in PXWEB_TABLES.items():
        params = {
            "lang": "sv",
            "outputFormat": "json-stat2",
            "valueCodes[Sektor]": "0",
            "valueCodes[Kon]": "1+2",
            "valueCodes[Yrke2012]": ",".join(CODES),
            "valueCodes[Tid]": "*",
            "valueCodes[ContentsCode]": "*",
        }
        params |= {f"valueCodes[{k}]": v for k, v in extra.items()}
        data = get(PXWEB.format(table=table), params=params).json()
        (OUT / "pxweb" / f"{table}.json").write_text(json.dumps(data, ensure_ascii=False, indent=1))


def jobtech() -> list[str]:
    data = get(JOBTECH, params={"query": JOBTECH_QUERY}).json()
    wanted = set(CODES) | {"1210"}  # 1210 is JobTech's merged code for 1211 and 1212
    concepts = [c for c in data["data"]["concepts"] if c["ssyk_code_2012"] in wanted]
    for c in concepts:  # keep the fixture small: first 8 occupation-names per code
        c["narrower"] = c["narrower"][:8]
    (OUT / "jobtech" / "ssyk_level_4.json").write_text(
        json.dumps({"data": {"concepts": concepts}}, ensure_ascii=False, indent=1)
    )
    uris = []
    for c in concepts:
        for n in c["narrower"]:
            for k in ("exact_match", "close_match"):
                uris += [m["esco_uri"] for m in n[k]]
    return sorted(set(uris))[:6]


def ilo(isco: set[str]) -> None:
    """ILO reserves reproduction rights, so keep only the real layout, levels, codes and titles;
    every text cell is synthetic."""
    src = openpyxl.load_workbook(io.BytesIO(get(ILO_URL).content), read_only=True).worksheets[0]
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = src.title
    keep = isco | {c[:n] for c in isco for n in (1, 2, 3)}  # unit groups and their parents
    for i, r in enumerate(src.iter_rows(values_only=True)):
        if i == 0:
            ws.append(list(r))
        elif r and (code := str(r[1]).strip()) in keep:
            ws.append([r[0], r[1], r[2],
                       f"Synthetic definition of ISCO-08 {code} (not ILO text).",
                       (f"Tasks include -\n(a)  synthetic task one of {code};\n"
                        f"(b)  synthetic task two of {code}."),
                       ("Examples of the occupations classified here:\n"
                        f"- Synthetic example A of {code}\n- Synthetic example B of {code}"),
                       None, None])
    out.save(OUT / "ilo" / "isco08_en.xlsx")


def esco(uris: list[str]) -> None:
    """One request per URI: ESCO returns every language at once. Keep en and sv only."""
    for i, uri in enumerate(uris):
        data = get(ESCO_URL, params={"uri": uri, "language": "en"}).json()
        keep = {"uri": data["uri"], "code": data.get("code")}
        for k in ("preferredLabel", "alternativeLabel", "description"):
            keep[k] = {lang: v for lang, v in (data.get(k) or {}).items() if lang in ("en", "sv")}
        (OUT / "esco" / f"occ{i}.json").write_text(json.dumps(keep, ensure_ascii=False, indent=1))


def mis2012() -> None:
    pages = [p.extract_text() or "" for p in PdfReader(io.BytesIO(get(MIS_URL).content)).pages]
    keep = {}
    for i, t in enumerate(pages):
        if any(re.search(rf"(?m)^\s*{c}\s", t) for c in ("1212", "4430", "5322")):
            keep[i] = t
    (OUT / "mis2012" / "pages.json").write_text(json.dumps(keep, ensure_ascii=False, indent=1))


def main() -> None:
    for d in ("scb", "pxweb", "jobtech", "ilo", "esco", "mis2012"):
        (OUT / d).mkdir(parents=True, exist_ok=True)
    isco = workbooks()
    pxweb()
    uris = jobtech()
    ilo(isco)
    esco(uris)
    mis2012()
    lines = ["# Test fixtures", "", "Real slices cut by `scripts/cut_fixtures.py` for codes "
             + ", ".join(CODES) + ".", "",
             ("Except `ilo/isco08_en.xlsx`: the ILO reserves reproduction rights, so it keeps "
              "only the real layout, codes and titles, and every text cell is synthetic."), "",
             "| URL | Fetched |", "|---|---|"]
    lines += [f"| {u} | {d} |" for u, d in sources]
    (OUT / "README.md").write_text("\n".join(lines) + "\n")
    print(f"wrote fixtures from {len(sources)} requests to {OUT}")


if __name__ == "__main__":
    main()
